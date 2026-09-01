# -*- coding: utf-8 -*-
"""Ajuste da tabela de comentarios das laminas de Small Caps (setembro/2026).

A carteira foi de 13 para 14 nomes e a tabela passou a invadir o rodape na
lamina XP (base 7,05 pol contra rodape em 6,89).

As alturas declaradas das linhas vem do template e nao correspondem ao
conteudo deste mes -- ECOR3 herdou 0,55 pol para um comentario de uma linha,
enquanto LWSA3 herdou 0,28 pol para um de tres. Como a altura renderizada e
max(declarada, natural), o excesso das linhas curtas e desperdicio puro.

Correcao: redimensionar cada linha do corpo pelo numero de linhas do
comentario e subir a tabela ate logo abaixo do titulo.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from pptx import Presentation
from pptx.util import Emu, Inches

# ~130 caracteres por linha na coluna de comentarios
POR_LINHA = 130
ALT = {1: 0.180, 2: 0.290, 3: 0.410}     # altura declarada por nº de linhas

ALVOS = [
    ('output/Carteira Top Small Caps - Setembro 2026.pptx', 3, 0.850),
    ('output/Carteira Small Caps Clear - Setembro 2026.pptx', 5, None),
]

for caminho, sl, novo_topo in ALVOS:
    prs = Presentation(caminho)
    slide = prs.slides[sl]
    tab = next(sh for sh in slide.shapes if sh.has_table
               and any('Coment' in c.text for c in sh.table.rows[0].cells))
    rodape = min(sh.top for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip().lower().startswith('fonte'))
    icom = next(i for i, c in enumerate(tab.table.rows[0].cells) if 'Coment' in c.text)

    antes = sum(r.height for r in tab.table.rows)
    for r in list(tab.table.rows)[1:]:
        n = min(3, max(1, -(-len(r.cells[icom].text.strip()) // POR_LINHA)))
        r.height = Inches(ALT[n])
    if novo_topo is not None:
        tab.top = Inches(novo_topo)
    depois = sum(r.height for r in tab.table.rows)
    tab.height = depois

    base = tab.top + depois
    print('%-44s soma %.3f -> %.3f | base=%.3f rodape=%.3f folga=%+.3f pol'
          % (caminho.split('/')[-1][:44], Emu(antes).inches, Emu(depois).inches,
             Emu(base).inches, Emu(rodape).inches, Emu(rodape - base).inches))
    prs.save(caminho)
