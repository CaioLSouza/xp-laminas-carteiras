# -*- coding: utf-8 -*-
"""Confere as laminas geradas contra a planilha e procura residuos do mes anterior.

    python verify.py --fechamento 2026-07 --xlsx "Charts Lamina Carteiras.xlsm"

Sai com codigo 1 se qualquer checagem falhar. As checagens existem porque cada
uma delas ja quebrou de verdade -- ver references/desafios.md.
"""
import sys, io, os, re, argparse, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from lxml import etree
import build as B

INGLES = ['top pick', 'top-pick', 'newsflow', 'backlog', 'valuation', 'ramp-up', 'sell-off',
          'preview', 'positioning', 'trigger', 'earnings', 'deal', 'deals', 'player', 'NIM',
          'ROE', 'ROEs', 'upside', 'underperform', 'outperform', 'guidance', 'capex',
          'premium', 'core', 'payout', 'turnaround', 'spread']
PERF = [r'(?:Alta|Queda|Baixa) de [\d,]+%', r'\d+(?:,\d+)? p\.p\.',
        r'(?:recuou|subiu|caiu|avançou|valorizou|desvalorizou) [\d,]+%',
        r'em \w+ \(-?[\d,]+%\)']
falhas = []


def chk(cond, msg):
    print(('  OK  ' if cond else '  XX  ') + msg)
    if not cond:
        falhas.append(msg)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--fechamento', required=True)
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--content', default='content.py')
    ap.add_argument('--out', default='output')
    ap.add_argument('--only', default=None)
    a = ap.parse_args()

    y, mo = (int(x) for x in a.fechamento.split('-'))
    prox = datetime.date(y + (mo == 12), mo % 12 + 1, 1)
    corte = datetime.datetime(prox.year, prox.month, 1) - datetime.timedelta(days=1)
    ant = corte.replace(day=1) - datetime.timedelta(days=1)
    mes_fech, mes_lam = B.MES_EXT[mo - 1], B.MES_EXT[prox.month - 1]
    ctx = dict(mes_ext=mes_lam, Mes=mes_lam.capitalize(), ano=prox.year)

    C = B.carrega_conteudo(a.content)
    planilha = B.Planilha(a.xlsx, corte)

    for key in (a.only.split(',') if a.only else list(B.DECKS)):
        cfg = B.DECKS[key]
        caminho = os.path.join(a.out, cfg['saida'].format(**ctx))
        print('=' * 72)
        print(os.path.basename(caminho))
        if not os.path.exists(caminho):
            chk(False, 'arquivo gerado')
            continue
        prs = Presentation(caminho)
        m = planilha.metrics(cfg['base100'])
        comp = planilha.rows(cfg['composicao'])
        ativos = planilha.rows(cfg['ativos'])
        textos = getattr(C, cfg['comentarios'])

        tabelas, txts = [], []
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    txts.append(sh.text_frame.text)
                if sh.has_table:
                    tabelas.append(sh.table)
                    txts += [c.text for r in sh.table.rows for c in r.cells]
        txt = '\n'.join(txts)

        # --- corte de dados: o erro mais caro do processo
        chk(corte.strftime('%d/%m/%Y') in txt, 'rodape com "Dados até %s"' % corte.strftime('%d/%m/%Y'))
        chk(ant.strftime('%d/%m/%Y') not in txt, 'sem resquicio de %s' % ant.strftime('%d/%m/%Y'))
        chk('%s %d' % (mes_lam.capitalize(), prox.year) in txt, 'capa rotulada com o mes da lamina')
        prox2 = B.MES[prox.month - 1]
        chk('%s-%02d' % (prox2, prox.year % 100) not in txt
            and '%s/%02d' % (prox2, prox.year % 100) not in txt,
            'nenhuma coluna do mes corrente parcial (%s)' % prox2)
        chk('Em %s,' % mes_fech in txt, 'texto de desempenho referente a %s' % mes_fech)

        r_c, r_i = m['months'][0][1][0], m['months'][0][1][1]
        chk(B.pc1(r_c, True) in txt and B.pc1(r_i, True) in txt,
            'retorno do mes na capa: %s vs %s' % (B.pc1(r_c, True), B.pc1(r_i, True)))

        # --- composicao (carteira nova) x planilha
        ctab = [t for t in tabelas if len(t.columns) == 9][0]
        got = [(ctab.cell(i, 5).text, ctab.cell(i, 6).text) for i in range(1, len(ctab.rows))]
        exp = [(r['Ticker'], B.pc1(r['Peso'], True)) for r in comp]
        chk(got == exp, 'composicao (%d nomes) confere com a planilha' % len(exp))
        if got != exp:
            print('     esperado:', exp, '\n     obtido  :', got)

        # --- desempenho por ativo = carteira do mes que fechou (abas _last)
        atab = [t for t in tabelas if len(t.columns) == 7
                and t.cell(0, 3).text.startswith('Data')][0]
        got = [(atab.cell(i, 1).text, atab.cell(i, 5).text) for i in range(1, len(atab.rows))]
        exp = [(r['Ticker'], B.pc1(r['Desempenho no mês'])) for r in ativos]
        chk(got == exp, 'desempenho por ativo (%d nomes, carteira de %s)' % (len(exp), mes_fech))
        if got != exp:
            print('     esperado:', exp, '\n     obtido  :', got)

        # --- performance: 12 meses terminando no fechamento
        ptab = [t for t in tabelas if t.cell(0, 1).text.strip().lower().startswith('desde')][0]
        hdr = [ptab.cell(0, i).text.strip() for i in range(len(ptab.columns))]
        i0 = next(i for i, h in enumerate(hdr) if h.lower().startswith('desde'))
        got = hdr[i0 + 3:i0 + 15]
        exp = [B.mlabel(yy, mm) for (yy, mm), _ in m['months']]
        chk(got == exp, 'cabecalho dos 12 meses: %s ... %s' % (got[0], got[-1]))
        ok = all(ptab.cell(j, i0 + 3 + k).text.strip() == B.pc1(m['months'][k][1][j - 1])
                 for j in range(1, cfg['nseries'] + 1) for k in range(12))
        chk(ok, 'valores mensais conferem com a serie diaria')
        ok = all(ptab.cell(j, c).text.strip() == B.pc1(v[j - 1])
                 for j in range(1, cfg['nseries'] + 1)
                 for c, v in ((i0, m['since']), (i0 + 1, m['ytd']), (i0 + 2, m['ltm'])))
        chk(ok, 'desde o inicio / ano / ult. 12M conferem (corte no fechamento)')

        # --- comentarios
        ttab = [t for t in tabelas if len(t.columns) == 7 and 'Coment' in t.cell(0, 5).text][0]
        tk = [ttab.cell(i, 1).text for i in range(1, len(ttab.rows))]
        chk(tk == [r['Ticker'] for r in comp], 'comentarios cobrem os %d nomes' % len(comp))
        chk(all(ttab.cell(i, 5).text.strip() == textos[tk[i - 1]].strip()
                for i in range(1, len(ttab.rows))), 'textos dos comentarios aplicados')
        links = [ttab.cell(i, 6).text_frame.paragraphs[0].runs[0].hyperlink.address
                 for i in range(1, len(ttab.rows))]
        chk(links == [B.TESE % t for t in tk], 'links da tese apontam para o ticker certo')

        largos = [(tk[i - 1], len(ttab.cell(i, 5).text)) for i in range(1, len(ttab.rows))
                  if len(ttab.cell(i, 5).text) > 385]
        chk(not largos, 'nenhum comentario estoura 3 linhas' + (' %s' % largos if largos else ''))
        ing = sorted({w for i in range(1, len(ttab.rows)) for w in INGLES
                      if re.search(r'\b%s\b' % re.escape(w), ttab.cell(i, 5).text, re.I)})
        chk(not ing, 'comentarios sem termos em ingles' + (' -> %s' % ing if ing else ''))
        perf = sorted({tk[i - 1] for i in range(1, len(ttab.rows)) for p in PERF
                       if re.search(p, ttab.cell(i, 5).text, re.I)})
        chk(not perf, 'comentarios sem numeros de desempenho do papel'
            + (' -> %s' % perf if perf else ''))

        # --- XML: endParaRPr tem de ser o ultimo filho de <a:p>.
        # Fora de ordem o PowerPoint apaga o texto e o LibreOffice nao mostra.
        ruins = 0
        for s in prs.slides:
            for sh in s.shapes:
                els = [sh.text_frame._txBody] if sh.has_text_frame else []
                if sh.has_table:
                    els += [c.text_frame._txBody for r in sh.table.rows for c in r.cells]
                for tx in els:
                    for para in tx.findall(B.A + 'p'):
                        kids = [etree.QName(c).localname for c in para]
                        if 'endParaRPr' in kids and kids.index('endParaRPr') != len(kids) - 1:
                            ruins += 1
        chk(ruins == 0, 'XML valido: endParaRPr sempre no fim do paragrafo'
            + ('' if not ruins else ' (%d fora de ordem)' % ruins))

    print('=' * 72)
    print('FALHAS:', len(falhas))
    for f in falhas:
        print(' -', f)
    sys.exit(1 if falhas else 0)


if __name__ == '__main__':
    main()
