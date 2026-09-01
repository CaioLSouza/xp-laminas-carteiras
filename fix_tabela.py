# -*- coding: utf-8 -*-
"""Ajuste de layout das capas de Small Caps (setembro/2026).

A carteira foi de 13 para 14 nomes. O `reposiciona()` do build.py dimensiona a
tabela por `shape.height`, mas quem manda no render e a SOMA das alturas de
linha -- e as duas divergem no template (2,58 vs 2,73 pol). Com 14 nomes a
soma passa 0,09 pol alem do topo do rodape.

Correcao: encolher a altura declarada das linhas do corpo (mantendo folga
sobre o minimo natural da fonte) e, na Clear, subir um pouco a tabela.
Alvo: 0,05 pol de folga entre a base da tabela e o topo do rodape.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from pptx import Presentation
from pptx.util import Emu, Inches

FOLGA = Inches(0.05)

# arquivo, slide da capa, nova altura de linha do corpo, novo topo (None = mantem)
ALVOS = [
    ('output/Carteira Top Small Caps - Setembro 2026.pptx', 0, 0.1510, None),
    ('output/Carteira Small Caps Clear - Setembro 2026.pptx', 2, 0.1560, 4.181),
]

for caminho, sl, alt_linha, novo_topo in ALVOS:
    prs = Presentation(caminho)
    slide = prs.slides[sl]
    tab = next(sh for sh in slide.shapes if sh.has_table and len(sh.table.rows) > 10)
    rodape = min(sh.top for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip().lower().startswith('fonte'))

    antes = sum(r.height for r in tab.table.rows)
    for r in list(tab.table.rows)[1:]:
        r.height = Inches(alt_linha)
    if novo_topo is not None:
        tab.top = Inches(novo_topo)
    depois = sum(r.height for r in tab.table.rows)
    tab.height = depois          # frame passa a refletir a altura real das linhas

    base = tab.top + depois
    folga = Emu(rodape - base).inches
    print('%-44s base=%.3f rodape=%.3f folga=%+.3f pol  %s'
          % (caminho.split('/')[-1][:44], Emu(base).inches, Emu(rodape).inches, folga,
             'OK' if base + FOLGA <= rodape + Inches(0.002) else 'AINDA ESTOURA'))
    print('     altura da tabela: %.4f -> %.4f pol' % (Emu(antes).inches, Emu(depois).inches))
    prs.save(caminho)
