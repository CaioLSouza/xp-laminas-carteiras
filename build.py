# -*- coding: utf-8 -*-
"""Monta as Laminas das Carteiras Recomendadas XP a partir dos templates.

    python build.py --fechamento 2026-07 --xlsx "Charts Lamina Carteiras.xlsm"

Gera as 4 laminas (Top Acoes, Top Dividendos, Top Small Caps e Small Caps Clear)
do mes seguinte ao fechamento. Todo dado e recalculado das series diarias
`* base 100` com corte no ultimo dia do mes de fechamento -- ver
references/desafios.md para o motivo.
"""
import os, re, sys, shutil, argparse, datetime, copy, importlib.util
import openpyxl
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.chart.data import CategoryChartData

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
CH = '{http://schemas.openxmlformats.org/drawingml/2006/chart}'
EPOCH = datetime.datetime(1899, 12, 30)
MES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
MES_EXT = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho',
           'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']

avisos = []


# ============================================================ configuracao
# Uma entrada por lamina. `shapes` mapeia (indice do slide, nome da shape).
# Os nomes vem do template e sobrevivem de um mes para o outro porque cada
# lamina nova e copia da anterior -- por isso os templates sao promovidos ao
# fim de cada rodada (ver SKILL.md).
DECKS = {
    'acoes': dict(
        template='Top Ações XP.pptx', saida='Top Ações XP - {Mes} {ano}.pptx',
        base100='Top Ações base 100', composicao='Composicao Top Ações',
        ativos='now_last_top_ações', metricas='Top Ações Portfolio Metrics',
        comentarios='ACOES', mudancas='MUDANCAS_ACOES', sub='SUB_ACOES', nseries=2,
        desempenho='Em {mes_ext}, a carteira teve retorno de {cart}, vs. {ibov} do Ibovespa.',
        shapes=dict(
            rotulo_mes=(1, 'Text Placeholder 6'), sub_capa=(1, 'Text Placeholder 28'),
            caixa_capa=(1, 'Espaço Reservado para Conteúdo 3'),
            texto_capa=(1, 'Espaço Reservado para Conteúdo 17'),
            tab_composicao=(1, 'Table 1'),
            sub_mud=(2, 'Text Placeholder 11'), caixa_mud=(2, 'Text Placeholder 3'),
            tab_perf=(3, 'Table 2'), tab_ativos=(3, 'Table 6'), tab_ind=(3, 'Table 36'),
            grafico=(3, 'Chart 9'), tab_coment=(4, 'Table 21'))),

    'divs': dict(
        template='Carteira Top Dividendos.pptx',
        saida='Carteira Top Dividendos - {Mes} {ano}.pptx',
        base100='Top Dividendos base 100', composicao='Composição Top Divs',
        ativos='now_last_top_divs', metricas='Top Div Portfolio Metrics',
        comentarios='DIVS', mudancas='MUDANCAS_DIVS', sub='SUB_DIVS', nseries=2,
        seg_map={'Defensivas Domésticas': 'Defensivas'},
        desempenho=('Em {mes_ext}, a carteira teve retorno de {cart}, vs. {ibov} do Ibovespa. '
                    'Desde o início da carteira, ela subiu {cart_ini} vs. {ibov_ini} do índice.'),
        shapes=dict(
            rotulo_mes=(1, 'Espaço Reservado para Texto 5'),
            sub_capa=(1, 'Espaço Reservado para Texto 14'),
            caixa_capa=(1, 'Espaço Reservado para Conteúdo 3'),
            texto_capa=(1, 'Espaço Reservado para Conteúdo 17'),
            tab_composicao=(1, 'Table 6'),
            sub_mud=(2, 'Text Placeholder 11'), caixa_mud=(2, 'Text Placeholder 3'),
            tab_ativos=(3, 'Table 4'), tab_ind=(3, 'Table 13'), grafico=(3, 'Chart 11'),
            tab_perf=(4, 'Table 11'), tab_coment=(5, 'Table 21'))),

    'small': dict(
        template='Carteira Top Small Caps.pptx',
        saida='Carteira Top Small Caps - {Mes} {ano}.pptx',
        base100='Top Small Caps base 100', composicao='Composição Top Smll',
        ativos='now_last_top_smll', metricas='Top SMLL Portfolio Metrics',
        comentarios='SMALL', mudancas='MUDANCAS_SMALL', sub='SUB_SMALL', nseries=3,
        seg_map={'Defensivas Domésticas': 'Defensivas', 'Bens Industriais': 'Industriais'},
        desempenho=('Em {mes_ext}, nossa carteira teve um desempenho de {cart}, vs. {ibov} do '
                    'Ibovespa. Desde o início, em maio de 2021, ela subiu {cart_ini}, vs. '
                    '{ibov_ini} do índice.'),
        shapes=dict(
            rotulo_mes=(1, 'Espaço Reservado para Texto 5'), sub_capa=(1, 'Text Placeholder 8'),
            caixa_capa=(1, 'Espaço Reservado para Conteúdo 3'),
            texto_capa=(1, 'Espaço Reservado para Conteúdo 17'),
            tab_composicao=(1, 'Table 9'),
            sub_mud=(2, 'Text Placeholder 4'), caixa_mud=(2, 'Text Placeholder 3'),
            tab_perf=(3, 'Table 2'), tab_ativos=(3, 'Table 12'), tab_ind=(3, 'Table 11'),
            grafico=(3, 'Chart 15'), tab_coment=(4, 'Table 21'))),

    'clear': dict(
        template='Carteira Small Caps Clear.pptx',
        saida='Carteira Small Caps Clear - {Mes} {ano}.pptx',
        base100='Top Small Caps base 100', composicao='Composição Top Smll',
        ativos='now_last_top_smll', metricas='Top SMLL Portfolio Metrics',
        comentarios='SMALL', mudancas='MUDANCAS_SMALL', sub='SUB_SMALL', nseries=3,
        seg_map={'Defensivas Domésticas': 'Defensivas', 'Bens Industriais': 'Industriais'},
        # a caixa da capa da Clear e 10pt e estoura quando o texto de mudancas
        # e longo; 9pt alinha com o paragrafo logo abaixo e com a lamina XP
        caixa_capa_pt=9,
        desempenho=('Em {mes_ext}, nossa carteira teve um desempenho de {cart}, vs. {ibov} do '
                    'Ibovespa. Desde o início, em maio de 2021, ela subiu {cart_ini}, vs. '
                    '{ibov_ini} do índice.'),
        shapes=dict(
            rotulo_mes=(3, 'Google Shape;132;p21'), sub_capa=(3, 'Google Shape;133;p21'),
            caixa_capa=(3, 'Google Shape;130;p21'), texto_capa=(3, 'Google Shape;136;p21'),
            tab_composicao=(3, 'Table 5'),
            sub_mud=(4, 'Google Shape;166;p23'), caixa_mud=(4, 'Google Shape;163;p23'),
            tab_perf=(5, 'Table 4'), tab_ativos=(5, 'Table 6'), tab_ind=(5, 'Table 9'),
            grafico=(5, 'Chart 1'), tab_coment=(6, 'Table 21'))),
}

NOMES = {'SLC Agricola': 'SLC Agrícola'}   # acentos que faltam na planilha
TESE = 'https://conteudos.xpi.com.br/acoes/%s'


# ============================================================ formatacao
def pc1(v, virgula=False):
    s = '%.1f%%' % (v * 100)
    return s.replace('.', ',') if virgula else s


def pc2(v):
    return '%.2f%%' % (v * 100)


def brl(v):
    return ('R$ %.2f' % v).replace('.', ',')


def mlabel(y, m):
    return '%s-%02d' % (MES[m - 1], y % 100)


def nome(v):
    return NOMES.get(v, v)


def as_dt(v):
    return v if isinstance(v, datetime.datetime) else EPOCH + datetime.timedelta(days=int(v))


# ============================================================ planilha
class Planilha:
    def __init__(self, caminho, corte):
        self.wb = openpyxl.load_workbook(caminho, data_only=True)
        self.corte = corte

    def rows(self, sheet):
        ws = self.wb[sheet]
        hdr = [c.value for c in ws[1]]
        return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)
                if r[0] is not None]

    def metrics(self, sheet):
        """Recalcula tudo das series diarias, com corte no fim do mes.

        As abas `performance_top_*` e `Portfolio Metrics` trazem o mes corrente
        parcial e contaminam ate o 'Desde o inicio'; a serie base 100 e a unica
        fonte confiavel. Ver references/desafios.md.
        """
        ws = self.wb[sheet]
        s = [(r[0], [float(x) for x in r[1:] if x is not None])
             for r in list(ws.iter_rows(values_only=True))[1:]
             if r[0] is not None and r[0] <= self.corte]
        if not s:
            raise SystemExit('ERRO: %s nao tem dados ate %s' % (sheet, self.corte.date()))
        if s[-1][0] < self.corte - datetime.timedelta(days=6):
            avisos.append('%s: ultimo dado em %s, bem antes do fechamento %s'
                          % (sheet, s[-1][0].date(), self.corte.date()))
        me = {(d.year, d.month): v for d, v in s}
        first, last = s[0][1], s[-1][1]
        n = len(first)
        rel = lambda a, b: [b[i] / a[i] - 1 for i in range(n)]
        ant = lambda y, m: (y - 1, 12) if m == 1 else (y, m - 1)

        alvo = (self.corte.year, self.corte.month)
        meses, ym = [], alvo
        for _ in range(12):
            py, pm = ant(*ym)
            if (py, pm) not in me or ym not in me:
                raise SystemExit('ERRO: %s sem dados para %s' % (sheet, mlabel(*ym)))
            meses.append((ym, rel(me[(py, pm)], me[ym])))
            ym = (py, pm)

        anos = {}
        for y in range(s[0][0].year, alvo[0]):
            base = me[(y - 1, 12)] if (y - 1, 12) in me else first
            if (y, 12) in me:
                anos[y] = rel(base, me[(y, 12)])
        base_ytd, base_ltm = (alvo[0] - 1, 12), (alvo[0] - 1, alvo[1])
        for b, rot in ((base_ytd, 'YTD'), (base_ltm, 'ultimos 12 meses')):
            if b not in me:
                raise SystemExit('ERRO: %s sem base para %s (%s)' % (sheet, rot, mlabel(*b)))
        return dict(series=s, since=rel(first, last), ytd=rel(me[base_ytd], last),
                    ltm=rel(me[base_ltm], last), months=meses, years=anos, ano=alvo[0])


# ============================================================ pptx: texto
def tf_of(o):
    return o.text_frame if hasattr(o, 'text_frame') else o


def set_text(obj, texto):
    """Troca o texto preservando a formatacao do primeiro run."""
    tf = tf_of(obj)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    p = tf.paragraphs[0]
    runs = p.runs
    if not runs:
        r = p.add_run()
        r.text = texto
        return r
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return runs[0]


def _append_run(p_el, r_el):
    """Insere o run ANTES de <a:endParaRPr>.

    Pelo schema DrawingML o endParaRPr tem de ser o ultimo filho de <a:p>.
    Fora de ordem, o PowerPoint apaga o texto em silencio -- e o LibreOffice
    renderiza normalmente, entao o QA visual nao pega. Ver desafios.md.
    """
    end = p_el.find(A + 'endParaRPr')
    end.addprevious(r_el) if end is not None else p_el.append(r_el)


def _modelos(paras):
    """Clona um run em negrito e um normal do template, para herdar fonte/cor."""
    b = n = None
    for p in paras:
        for r in p.runs:
            if r.font.bold and b is None:
                b = copy.deepcopy(r._r)
            if not r.font.bold and n is None:
                n = copy.deepcopy(r._r)
    if b is None and n is None:
        raise ValueError('shape sem run para servir de modelo de formatacao')
    return (b if b is not None else copy.deepcopy(n),
            n if n is not None else copy.deepcopy(b))


def set_rich_para(para, trechos, modelo_paras=None):
    """Reescreve um paragrafo com [(texto, negrito), ...]."""
    src_b, src_n = _modelos(modelo_paras or [para])
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    for txt, bold in trechos:
        r = copy.deepcopy(src_b if bold else src_n)
        for t in r.findall(A + 't'):
            r.remove(t)
        t = r.makeelement(A + 't', {})
        t.text = txt
        if txt != txt.strip():
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r.append(t)
        _append_run(para._p, r)
    return para


def set_rich(obj, trechos):
    tf = tf_of(obj)
    modelos = list(tf.paragraphs)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    return set_rich_para(tf.paragraphs[0], trechos, modelos)


def set_size(obj, pt):
    for para in tf_of(obj).paragraphs:
        for r in para.runs:
            r.font.size = Pt(pt)


def substituir(prs, alvos, novo):
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    for velho in alvos:
                        if velho in r.text:
                            r.text = r.text.replace(velho, novo)


def datas_dinamicas(prs, texto):
    """Atualiza o cache dos campos <a:fld type=datetime*> (eles se atualizam
    sozinhos ao abrir, mas o cache e o que aparece em pre-visualizacoes)."""
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for fld in p._p.findall(A + 'fld'):
                    if (fld.get('type') or '').startswith('datetime'):
                        for t in fld.findall(A + 't'):
                            t.text = texto


# ============================================================ pptx: tabelas
def achar(slide, nome, ctx=''):
    for sh in slide.shapes:
        if sh.name == nome:
            return sh
    raise SystemExit("ERRO: shape '%s' nao existe %s.\n"
                     "  O template mudou? Rode scripts/inspect.py para listar as shapes "
                     "e atualize DECKS em build.py." % (nome, ctx))


def trs(table):
    return table._tbl.findall(A + 'tr')


def set_rowcount(table, n, modelo=2):
    modelo = min(modelo, len(trs(table)) - 1)
    while len(trs(table)) - 1 > n:
        tr = trs(table)[-1]
        tr.getparent().remove(tr)
    while len(trs(table)) - 1 < n:
        trs(table)[-1].addnext(copy.deepcopy(trs(table)[modelo]))


def normalize_body(table, ref=1):
    """Uniformiza corpo e bordas do corpo pela linha de referencia.

    Celulas que eram continuacao de mesclagem nao tem run formatado; escrever
    nelas direto produziria texto no tamanho padrao (18pt) com borda visivel.
    """
    ncols = len(table.columns)
    refs = [copy.deepcopy(table.cell(ref, c)._tc) for c in range(ncols)]
    for ri in range(1, len(table.rows)):
        for ci in range(ncols):
            tc = table.cell(ri, ci)._tc
            for tag in ('txBody', 'tcPr'):
                velho, novo = tc.find(A + tag), refs[ci].find(A + tag)
                if velho is not None and novo is not None:
                    tc.replace(velho, copy.deepcopy(novo))


def unmerge(table, cols):
    for ri in range(1, len(table.rows)):
        for ci in cols:
            tc = table.cell(ri, ci)._tc
            tc.attrib.pop('vMerge', None)
            tc.attrib.pop('rowSpan', None)


def vmerge(table, ci, r0, span):
    if span <= 1:
        return
    table.cell(r0, ci)._tc.set('rowSpan', str(span))
    for r in range(r0 + 1, r0 + span):
        table.cell(r, ci)._tc.set('vMerge', '1')


def blocos(valores):
    out = []
    for i, v in enumerate(valores):
        if out and out[-1][2] == v:
            out[-1][1] += 1
        else:
            out.append([i, 1, v])
    return out


def set_link(cell, url):
    cell.text_frame.paragraphs[0].runs[0].hyperlink.address = url


def reposiciona(shape, slide, linhas_antes, linhas_depois):
    """Reajusta altura da tabela e sobe se ela passar a invadir o rodape.

    A carteira muda de tamanho de um mes para o outro; sem isso a tabela de
    composicao encosta na fonte quando entra um nome novo.
    """
    if linhas_depois == linhas_antes:
        return
    por_linha = shape.height / linhas_antes
    nova = int(por_linha * linhas_depois)
    shape.height = nova
    rodape = [sh.top for sh in slide.shapes
              if sh.has_text_frame and sh.text_frame.text.strip().lower().startswith('fonte')]
    if rodape:
        limite = min(rodape) - Inches(0.06)
        if shape.top + nova > limite:
            shape.top = limite - nova


# ============================================================ blocos da lamina
def fill_composicao(shape, comp, seg_map=None, slide=None):
    table = shape.table
    seg_map = seg_map or {}
    n, antes = len(comp), len(table.rows) - 1
    altura = trs(table)[1].get('h')
    set_rowcount(table, n)
    normalize_body(table)
    if altura:
        for tr in trs(table)[1:]:
            tr.set('h', altura)
    unmerge(table, [0, 1, 2, 3])
    segs = [seg_map.get(r['Segmento'], r['Segmento']) for r in comp]
    sets = [seg_map.get(r['Setor'], r['Setor']) for r in comp]
    for i, r in enumerate(comp):
        ri = i + 1
        for ci, v in enumerate([segs[i], sets[i],
                                pc1(r['Peso do setor (Ibovespa)'], True),
                                pc1(r['Peso do setor (Carteira)'], True),
                                nome(r['Companhia']), r['Ticker'],
                                pc1(r['Peso'], True), r['Rating'], brl(r['Preço-Alvo'])]):
            set_text(table.cell(ri, ci), v)
    for i, sz, _ in blocos(segs):
        vmerge(table, 0, i + 1, sz)
    for i, sz, _ in blocos(['%s|%s' % (segs[k], sets[k]) for k in range(n)]):
        for ci in (1, 2, 3):
            vmerge(table, ci, i + 1, sz)
    if slide is not None:
        reposiciona(shape, slide, antes + 1, n + 1)


def fill_ativos(table, dados):
    """Figura 'Desempenho de cada ativo': carteira do mes que fechou.

    Usa as abas `*_last` -- e a carteira efetivamente detida no mes, entao
    lista nomes que ja sairam. Nao confundir com a composicao nova da capa.
    """
    set_rowcount(table, len(dados))
    for i, r in enumerate(dados):
        ri = i + 1
        for ci, v in enumerate([nome(r['Companhia']), r['Ticker'], pc2(r['Peso']),
                                mlabel(as_dt(r['Data de entrada']).year,
                                       as_dt(r['Data de entrada']).month),
                                pc1(r['Desempenho desde entrada']),
                                pc1(r['Desempenho no mês']), pc1(r['Desempenho YTD'])]):
            set_text(table.cell(ri, ci), v)


def fill_perf(table, m, nlinhas):
    """Tabela de performance: desde o inicio | ano | ult. 12M | 12 meses | anos.

    As colunas sao localizadas pelo cabecalho, nao por indice fixo, para o
    layout poder variar entre as laminas.
    """
    hdr = [table.cell(0, i).text.strip() for i in range(len(table.columns))]
    try:
        i0 = next(i for i, h in enumerate(hdr) if h.lower().startswith('desde'))
    except StopIteration:
        raise SystemExit("ERRO: tabela de performance sem coluna 'Desde o início'")
    m0 = i0 + 3
    set_text(table.cell(0, i0 + 1), str(m['ano']))
    for k, ((y, mo), vals) in enumerate(m['months']):
        set_text(table.cell(0, m0 + k), mlabel(y, mo))
        for j in range(1, nlinhas + 1):
            set_text(table.cell(j, m0 + k), pc1(vals[j - 1]))
    for j in range(1, nlinhas + 1):
        for ci, serie in ((i0, m['since']), (i0 + 1, m['ytd']), (i0 + 2, m['ltm'])):
            set_text(table.cell(j, ci), pc1(serie[j - 1]))
    for i in range(m0 + 12, len(hdr)):
        if re.fullmatch(r'\d{4}', hdr[i]):
            y = int(hdr[i])
            if y in m['years']:
                for j in range(1, nlinhas + 1):
                    set_text(table.cell(j, i), pc1(m['years'][y][j - 1]))
            else:
                avisos.append('coluna anual %d sem dado na serie -- confira a tabela' % y)
    anuais = {int(h) for h in hdr[m0 + 12:] if re.fullmatch(r'\d{4}', h)}
    if anuais and (m['ano'] - 1) in m['years'] and (m['ano'] - 1) not in anuais:
        avisos.append('virada de ano: %d fechou e a tabela de performance nao tem coluna para '
                      'ele -- abra o template e ajuste as colunas anuais' % (m['ano'] - 1))


def fill_indicadores(table, planilha, sheet):
    """Sharpe / Volatilidade / Beta saem da planilha como estao.

    Sao calculados pela metodologia da casa; recalcular daria numeros bem
    diferentes. O pregao a mais que essas abas carregam e imaterial numa
    serie de centenas de dias. Ver desafios.md.
    """
    ws = planilha.wb[sheet]
    for ri, linha in enumerate([[c.value for c in r] for r in ws.iter_rows(min_row=2)], start=1):
        set_text(table.cell(ri, 0), linha[0] + ' ')
        for ci, v in enumerate(linha[1:], start=1):
            if v is None:
                continue
            set_text(table.cell(ri, ci), pc2(v) if linha[0] == 'Volatilidade' else '%.2f' % v)


def _nomes_series(chart):
    out = []
    for ser in chart._chartSpace.iter(CH + 'ser'):
        tx = ser.find(CH + 'tx')
        v = tx.find(CH + 'v') if tx is not None else None
        if v is None and tx is not None:
            ref = tx.find(CH + 'strRef')
            cache = ref.find(CH + 'strCache') if ref is not None else None
            pt = cache.find(CH + 'pt') if cache is not None else None
            v = pt.find(CH + 'v') if pt is not None else None
        out.append(v.text if v is not None else '')
    return out


def _cache(el, valores):
    for pt in el.findall(CH + 'pt'):
        el.remove(pt)
    el.find(CH + 'ptCount').set('val', str(len(valores)))
    for i, v in enumerate(valores):
        pt = el.makeelement(CH + 'pt', {'idx': str(i)})
        cv = pt.makeelement(CH + 'v', {})
        cv.text = ('%d' % v) if isinstance(v, int) else ('%.10g' % v)
        pt.append(cv)
        el.append(pt)


def _tem_workbook(chart):
    """True se o grafico tem planilha embutida.

    Quando o vinculo e externo (lamina Clear), python-pptx levanta ValueError
    em vez de devolver None -- o relacionamento existe, mas em target-mode
    external.
    """
    try:
        return chart.part.chart_workbook.xlsx_part is not None
    except ValueError:
        return False


def fill_chart(chart, series):
    """Atualiza o grafico base 100.

    Se o grafico tem planilha embutida, replace_data() e o caminho normal.
    A lamina Clear aponta para um .xlsm externo (caminho de rede) e nao tem
    embedding -- ali reescrevemos o cache no XML, preservando o vinculo e os
    nomes das series.
    """
    if _tem_workbook(chart):
        cd = CategoryChartData()
        cd.categories = [d.date() for d, _ in series]
        cd.categories.number_format = r'yyyy\-mm\-dd'
        for i, nm in enumerate(_nomes_series(chart)):
            cd.add_series(nm, [v[i] for _, v in series], number_format='General')
        chart.replace_data(cd)
        return 'embutida'
    dias = [(d - EPOCH).days for d, _ in series]
    for i, ser in enumerate(chart._chartSpace.iter(CH + 'ser')):
        _cache(ser.find(CH + 'cat').find(CH + 'numRef').find(CH + 'numCache'), dias)
        _cache(ser.find(CH + 'val').find(CH + 'numRef').find(CH + 'numCache'),
               [v[i] for _, v in series])
    return 'vinculada'


def fill_comentarios(table, comp, textos):
    faltam = [r['Ticker'] for r in comp if r['Ticker'] not in textos]
    if faltam:
        raise SystemExit('ERRO: sem comentario para %s. Todo nome da carteira precisa de '
                         'uma entrada no arquivo de conteudo.' % ', '.join(faltam))
    set_rowcount(table, len(comp))
    for i, r in enumerate(comp):
        ri, tk = i + 1, r['Ticker']
        for ci, v in enumerate([nome(r['Companhia']), tk, pc1(r['Peso'], True), r['Rating'],
                                brl(r['Preço-Alvo']), textos[tk], 'Clique aqui']):
            set_text(table.cell(ri, ci), v)
        set_link(table.cell(ri, 6), TESE % tk)


# ============================================================ montagem
def build(key, cfg, planilha, C, ctx, templates, saida):
    origem = os.path.join(templates, cfg['template'])
    if not os.path.exists(origem):
        raise SystemExit('ERRO: template ausente: %s' % origem)
    destino = os.path.join(saida, cfg['saida'].format(**ctx))
    shutil.copy(origem, destino)
    prs = Presentation(destino)
    m = planilha.metrics(cfg['base100'])
    comp = planilha.rows(cfg['composicao'])
    ativos = planilha.rows(cfg['ativos'])
    sh = lambda k: achar(prs.slides[cfg['shapes'][k][0] - 1], cfg['shapes'][k][1],
                         'em %s (slide %d)' % (key, cfg['shapes'][k][0]))

    mud = getattr(C, cfg['mudancas'])
    sub = getattr(C, cfg['sub'])
    set_text(sh('rotulo_mes'), '%s %d' % (ctx['Mes'], ctx['ano']))
    set_text(sh('sub_capa'), sub)
    set_text(sh('sub_mud'), sub)
    caixa = sh('caixa_capa')
    set_rich(caixa, mud)
    if cfg.get('caixa_capa_pt'):
        set_size(caixa, cfg['caixa_capa_pt'])
    set_rich(sh('caixa_mud'), mud)

    tf = sh('texto_capa').text_frame
    set_rich_para(tf.paragraphs[-1], [
        ('Desempenho. ', True),
        (cfg['desempenho'].format(
            mes_ext=ctx['mes_fech_ext'],
            cart=pc1(m['months'][0][1][0], True), ibov=pc1(m['months'][0][1][1], True),
            cart_ini=pc1(m['since'][0], True), ibov_ini=pc1(m['since'][1], True)), False)],
        modelo_paras=list(tf.paragraphs))

    tab_comp = sh('tab_composicao')
    fill_composicao(tab_comp, comp, cfg.get('seg_map'),
                    slide=prs.slides[cfg['shapes']['tab_composicao'][0] - 1])
    fill_perf(sh('tab_perf').table, m, cfg['nseries'])
    fill_ativos(sh('tab_ativos').table, ativos)
    fill_indicadores(sh('tab_ind').table, planilha, cfg['metricas'])
    modo = fill_chart(sh('grafico').chart, m['series'])
    fill_comentarios(sh('tab_coment').table, comp, getattr(C, cfg['comentarios']))

    substituir(prs, ['Dados até %s' % ctx['data_ant'], 'Dados de %s' % ctx['data_ant']],
               'Dados até %s' % ctx['data_corte'])
    datas_dinamicas(prs, ctx['data_pub'])
    prs.save(destino)
    print('  -> %-46s (%d nomes, grafico %s)' % (os.path.basename(destino), len(comp), modo))
    return destino


def carrega_conteudo(caminho):
    if not os.path.exists(caminho):
        raise SystemExit('ERRO: arquivo de conteudo nao encontrado: %s\n'
                         '  Escreva-o a partir de scripts/content_exemplo.py.' % caminho)
    spec = importlib.util.spec_from_file_location('conteudo', caminho)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--fechamento', required=True, metavar='AAAA-MM',
                    help='mes que fechou (a lamina sai rotulada com o mes seguinte)')
    ap.add_argument('--xlsx', required=True, help='Charts Lamina Carteiras.xlsm')
    ap.add_argument('--content', default='content.py', help='comentarios e textos do mes')
    ap.add_argument('--templates', default=None, help='pasta dos templates (padrao: da skill)')
    ap.add_argument('--out', default='output', help='pasta de saida')
    ap.add_argument('--data-pub', default=None, metavar='AAAA-MM-DD',
                    help='data de publicacao (padrao: hoje)')
    ap.add_argument('--only', default=None,
                    help='subconjunto, ex.: acoes,clear (padrao: todas)')
    a = ap.parse_args()

    y, mo = (int(x) for x in a.fechamento.split('-'))
    prox = datetime.date(y + (mo == 12), mo % 12 + 1, 1)
    corte = datetime.datetime(prox.year, prox.month, 1) - datetime.timedelta(days=1)
    ant = corte.replace(day=1) - datetime.timedelta(days=1)
    pub = (datetime.date.fromisoformat(a.data_pub) if a.data_pub else datetime.date.today())
    ctx = dict(mes_ext=MES_EXT[prox.month - 1], Mes=MES_EXT[prox.month - 1].capitalize(),
               ano=prox.year,
               mes_fech_ext=MES_EXT[mo - 1],
               data_corte=corte.strftime('%d/%m/%Y'), data_ant=ant.strftime('%d/%m/%Y'),
               data_pub='%d de %s de %d' % (pub.day, MES_EXT[pub.month - 1], pub.year))

    templates = a.templates or os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                            '..', 'templates')
    os.makedirs(a.out, exist_ok=True)
    C = carrega_conteudo(a.content)
    planilha = Planilha(a.xlsx, corte)
    escolhidas = a.only.split(',') if a.only else list(DECKS)

    print('Fechamento %s | laminas de %s/%d | dados ate %s'
          % (a.fechamento, ctx['mes_ext'], ctx['ano'], ctx['data_corte']))
    for key in escolhidas:
        if key not in DECKS:
            raise SystemExit('ERRO: lamina desconhecida %r (use %s)' % (key, ', '.join(DECKS)))
        build(key, DECKS[key], planilha, C, ctx, templates, a.out)

    if avisos:
        print('\nAVISOS (confira antes de publicar):')
        for x in dict.fromkeys(avisos):
            print('  ! ' + x)


if __name__ == '__main__':
    main()
